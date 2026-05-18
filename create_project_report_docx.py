from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PROJECT_DIR = Path(__file__).resolve().parent
REPORT_PATH = PROJECT_DIR / "Credit_Card_Fraud_Detection_Project_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(hdr[index], header, bold=True)
        set_cell_shading(hdr[index], "E2E8F0")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    document.add_paragraph()
    return table


def add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(31, 78, 121)
    return heading


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Credit Card Fraud Detection System")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Using Random Forest Algorithm and Streamlit")
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(14)
    subtitle_run.italic = True

    doc.add_paragraph()
    add_table(
        doc,
        ["Item", "Details"],
        [
            ["Project Type", "Final Year Mini Project"],
            ["Algorithm", "Random Forest Classifier"],
            ["Frontend", "Streamlit"],
            ["Language", "Python"],
            ["Output", "Safe Transaction or Fraud Transaction Detected"],
        ],
    )

    doc.add_page_break()

    add_heading(doc, "1. Abstract")
    doc.add_paragraph(
        "Credit card fraud is a major issue in digital payment systems. This project presents a mini "
        "machine learning system that predicts whether a credit card transaction is safe or fraudulent. "
        "The system uses a Random Forest Classifier and displays the output through a Streamlit web "
        "application. The user manually enters transaction details, and the trained model predicts the "
        "transaction class along with a fraud probability score."
    )

    add_heading(doc, "2. Introduction")
    doc.add_paragraph(
        "Credit card transactions are widely used in online shopping, banking, travel booking, and point-of-sale "
        "payments. With the increase in digital transactions, fraud detection has become an important area in "
        "financial security. Machine learning can help identify suspicious patterns from transaction data."
    )

    add_heading(doc, "3. Problem Statement")
    doc.add_paragraph(
        "The aim of this project is to develop a credit card fraud detection system that accepts manually entered "
        "transaction details and classifies the transaction as safe or fraudulent using the Random Forest algorithm."
    )

    add_heading(doc, "4. Objectives")
    add_bullets(
        doc,
        [
            "To build a machine learning based fraud detection model.",
            "To use Random Forest Classifier for transaction classification.",
            "To create a Streamlit application for manual user input.",
            "To display fraud or safe transaction output clearly.",
            "To make the project simple and suitable for academic demonstration.",
        ],
    )

    add_heading(doc, "5. Proposed System")
    doc.add_paragraph(
        "The proposed system uses transaction details such as amount, transaction hour, failed attempts, distance "
        "from home, account age, card presence, international transaction status, transaction type, and merchant "
        "category. These details are passed to the trained Random Forest model. The model predicts whether the "
        "transaction is safe or fraudulent."
    )

    add_heading(doc, "6. Hardware and Software Requirements")
    add_table(
        doc,
        ["Requirement", "Details"],
        [
            ["Processor", "Intel i3 or above"],
            ["RAM", "4 GB or above"],
            ["Storage", "500 MB free space"],
            ["Operating System", "Windows / Linux / macOS"],
            ["Python", "Python 3.10 or above"],
        ],
    )
    add_table(
        doc,
        ["Software", "Purpose"],
        [
            ["Python", "Programming language"],
            ["Streamlit", "Web interface"],
            ["scikit-learn", "Machine learning algorithm"],
            ["pandas", "Data processing"],
            ["numpy", "Numerical operations"],
            ["joblib", "Model saving and loading"],
        ],
    )

    add_heading(doc, "7. Dataset Description")
    doc.add_paragraph(
        "The dataset is generated inside the training script, so the project does not depend on an external dataset. "
        "The target column is is_fraud, where 0 represents a safe transaction and 1 represents a fraud transaction."
    )
    add_table(
        doc,
        ["Feature", "Description"],
        [
            ["amount", "Transaction amount"],
            ["hour", "Hour of transaction from 0 to 23"],
            ["previous_transactions_24h", "Number of previous transactions in the last 24 hours"],
            ["failed_attempts", "Failed login or payment attempts"],
            ["distance_from_home_km", "Distance from user's normal location"],
            ["account_age_days", "Age of account in days"],
            ["card_present", "Whether physical card is present"],
            ["international_transaction", "Whether transaction is international"],
            ["transaction_type", "POS, Online, ATM, or Transfer"],
            ["merchant_category", "Merchant type"],
        ],
    )

    add_heading(doc, "8. Algorithm: Random Forest")
    doc.add_paragraph(
        "Random Forest is a supervised machine learning algorithm that creates many decision trees and combines "
        "their outputs to make a final prediction. It is suitable for classification problems and works well with "
        "tabular data. In this project, Random Forest performs binary classification: safe transaction or fraud "
        "transaction."
    )
    add_bullets(
        doc,
        [
            "It reduces overfitting compared to a single decision tree.",
            "It handles complex transaction patterns.",
            "It gives good performance for classification problems.",
            "It is easy to explain in an academic project.",
        ],
    )

    add_heading(doc, "9. System Architecture")
    add_bullets(
        doc,
        [
            "Generate transaction dataset.",
            "Preprocess numerical and categorical features.",
            "Train the Random Forest Classifier.",
            "Save the trained model.",
            "Load the model in Streamlit.",
            "Accept manual transaction input.",
            "Display prediction result and fraud probability.",
        ],
    )

    add_heading(doc, "10. Module Description")
    add_table(
        doc,
        ["Module", "Description"],
        [
            ["Data Generation", "Creates sample credit card transaction records."],
            ["Preprocessing", "Scales numerical data and encodes categorical data."],
            ["Model Training", "Trains the Random Forest Classifier."],
            ["Model Saving", "Stores the trained model as a pickle file."],
            ["Streamlit App", "Accepts manual user input and displays output."],
        ],
    )

    add_heading(doc, "11. Implementation")
    doc.add_paragraph(
        "The project mainly contains two Python files. The train_model.py file creates the dataset, trains the "
        "Random Forest model, evaluates it, and saves the model. The app.py file loads the trained model, accepts "
        "manual input, predicts fraud probability, and displays the final result."
    )

    add_heading(doc, "12. Testing")
    add_table(
        doc,
        ["Test Case", "Input Type", "Expected Output", "Status"],
        [
            ["TC-01", "Application launch", "Streamlit app opens", "Pass"],
            ["TC-02", "Safe transaction details", "Safe Transaction", "Pass"],
            ["TC-03", "Suspicious transaction details", "Fraud Transaction Detected", "Pass"],
            ["TC-04", "Run training script", "Model file generated", "Pass"],
        ],
    )

    add_heading(doc, "13. Results")
    doc.add_paragraph(
        "The Random Forest model was trained successfully. The Streamlit application accepts manual input and "
        "shows the prediction result. During testing, the safe example was predicted as safe, and the fraud example "
        "was predicted as fraudulent."
    )

    add_heading(doc, "14. Advantages")
    add_bullets(
        doc,
        [
            "Simple and easy to use.",
            "Uses Random Forest algorithm.",
            "Accepts manual transaction input.",
            "Displays fraud probability.",
            "Does not require external dataset download.",
        ],
    )

    add_heading(doc, "15. Limitations")
    add_bullets(
        doc,
        [
            "Dataset is generated for academic demonstration.",
            "The system is not connected to a real banking server.",
            "It does not process live payment gateway transactions.",
            "Real-world deployment would require more data and security controls.",
        ],
    )

    add_heading(doc, "16. Future Enhancements")
    add_bullets(
        doc,
        [
            "Use a real credit card transaction dataset.",
            "Add login and admin dashboard.",
            "Store transaction history in a database.",
            "Add email or SMS fraud alerts.",
            "Deploy the system online.",
            "Compare Random Forest with other machine learning algorithms.",
        ],
    )

    add_heading(doc, "17. Conclusion")
    doc.add_paragraph(
        "This project successfully implements a credit card fraud detection system using Random Forest. The system "
        "accepts manually entered transaction details and predicts whether the transaction is safe or fraudulent. "
        "The Streamlit interface makes the project easy to demonstrate and suitable for final-year mini project "
        "submission."
    )

    add_heading(doc, "18. References")
    add_bullets(
        doc,
        [
            "scikit-learn documentation",
            "Streamlit documentation",
            "pandas documentation",
            "Machine learning classification concepts",
        ],
    )

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    build_document()
